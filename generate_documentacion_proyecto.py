from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(r"C:\Users\pemip\Desktop\Proyecto-ia-dental")
OUT = ROOT / "DOCUMENTACION_PROYECTO.docx"
IMG_DIR = ROOT / "capturas"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_after = Pt(6)
    return p


def paragraph(doc, text, align=None, bold=False, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def code_block(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = True
    hdr = t.rows[0]
    repeat_header(hdr)
    for idx, header in enumerate(headers):
        set_cell_text(hdr.cells[idx], header, bold=True)
        shade_cell(hdr.cells[idx], "D9EAF7")
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    doc.add_paragraph()


def add_image(doc, path, caption, width_inches=6.1):
    if not path.exists():
        paragraph(doc, f"[Imagen no disponible: {path.name}]", italic=True)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))

    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap.runs:
        cap.runs[0].italic = True
    cap.paragraph_format.space_after = Pt(10)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    # Portada
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Proyecto IA Dental")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Documentación técnica y funcional")
    r.font.name = "Arial"
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Fecha de elaboración: {date(2026, 5, 23).strftime('%d/%m/%Y')}").italic = True

    paragraph(
        doc,
        "Documento preparado para revisión, carga en Google Drive y uso como guía de implementación, despliegue y mantenimiento.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    doc.add_paragraph()

    heading(doc, "1. Resumen ejecutivo", 1)
    paragraph(doc, "Proyecto IA Dental es una plataforma web para análisis de radiografías dentales con apoyo de inteligencia artificial. El sistema combina un frontend en Vue 3, un backend en FastAPI con GraphQL, almacenamiento en MongoDB y un modelo YOLO para detección de hallazgos clínicos.")
    paragraph(doc, "La aplicación está organizada para cubrir tres flujos principales: autenticación de usuarios, carga y análisis de imágenes, y consulta del historial clínico con métricas resumidas en un dashboard responsivo.")

    heading(doc, "2. Objetivo y alcance", 1)
    bullet(doc, "Permitir el registro e inicio de sesión de profesionales.")
    bullet(doc, "Recibir radiografías en base64 desde el frontend y procesarlas en el backend.")
    bullet(doc, "Ejecutar inferencia con el modelo YOLO y almacenar resultados en MongoDB.")
    bullet(doc, "Mostrar resultados clínicos, histórico y métricas resumidas en la interfaz.")
    bullet(doc, "Mantener una arquitectura extensible con eventos, soporte Kafka y registros estructurados.")

    heading(doc, "3. Tecnologías principales", 1)
    table(
        doc,
        ["Tecnología", "Uso en el proyecto", "Observaciones"],
        [
            ["Vue 3", "Interfaz de usuario", "SPA con navegación por rutas y composables."],
            ["Vite", "Servidor de desarrollo y build", "Usa variables VITE_ en tiempo de compilación."],
            ["Tailwind CSS + daisyUI", "Estilos y componentes visuales", "El dashboard usa flexbox, no CSS Grid."],
            ["FastAPI", "API backend", "Expone health check y GraphQL."],
            ["Strawberry GraphQL", "Capa GraphQL", "Query, mutation y tipado compartido."],
            ["MongoDB + Motor", "Persistencia", "Usuarios, análisis y dead letters."],
            ["Ultralytics YOLO", "Inferencia IA", "Carga el peso best.pt configurado."],
            ["Kafka", "Eventos asíncronos", "Opcional, con fallback a logs."],
            ["Pillow", "Normalización de imágenes", "Convierte imágenes a PNG temporal cuando hace falta."],
        ],
    )

    heading(doc, "4. Estructura del repositorio", 1)
    table(
        doc,
        ["Ruta", "Propósito"],
        [
            ["backend/", "API, lógica de negocio, persistencia, eventos e inferencia."],
            ["frontend/", "Aplicación web en Vue 3 para autenticación, carga y dashboard."],
            ["entrenamiento ia/", "Dataset, scripts, pesos del modelo y pruebas del entrenamiento."],
            ["docker/", "Composición de infraestructura y utilidades para servicios auxiliares."],
            ["capturas/", "Evidencias visuales de entrenamiento, validación y predicción."],
            ["start_aplicacion.ps1 / start_aplicaion.sh", "Arranque interactivo del sistema completo."],
        ],
    )

    heading(doc, "5. Backend", 1)
    paragraph(doc, "El backend está implementado como una aplicación FastAPI con GraphQL y arranque centralizado en `backend/main.py`, que instancia `src.app.create_app()`.")
    paragraph(doc, "Durante el startup se conectan MongoDB, repositorios, servicios de autenticación y análisis, el cargador del modelo YOLO y, opcionalmente, los productores y consumidores de eventos.")

    heading(doc, "5.1. Puntos de entrada y flujo de arranque", 2)
    table(
        doc,
        ["Archivo", "Función"],
        [
            ["backend/main.py", "Expone la variable `app` para Uvicorn."],
            ["backend/src/app.py", "Construye FastAPI, monta GraphQL, CORS, static files y lifecycle startup/shutdown."],
            ["backend/src/config/settings.py", "Resuelve configuración por entorno y rutas absolutas."],
        ],
    )
    numbered(doc, "FastAPI arranca y configura CORS y GraphQL.")
    numbered(doc, "Se crea el directorio de uploads y se conecta MongoDB.")
    numbered(doc, "Se crean índices en `analyses`, `users` y `dead_letter_events`.")
    numbered(doc, "Se carga el modelo YOLO si `model_warmup_on_startup` está activo.")
    numbered(doc, "Se inicializan el publicador de eventos y los servicios de negocio.")

    heading(doc, "5.2. Operaciones GraphQL", 2)
    table(
        doc,
        ["Operación", "Tipo", "Acceso", "Descripción"],
        [
            ["registerUser", "Mutation", "Pública", "Crea un usuario nuevo."],
            ["loginUser", "Mutation", "Pública", "Devuelve JWT y datos del usuario."],
            ["refreshToken", "Mutation", "Autenticado", "Renueva el token activo."],
            ["logoutUser", "Mutation", "Autenticado", "Publica el evento de cierre de sesión."],
            ["uploadRadiography", "Mutation", "Autenticado", "Guarda la imagen, ejecuta inferencia y persiste el análisis."],
            ["getAnalysisById", "Query", "Autenticado", "Consulta un análisis propio o, si es admin, cualquier análisis."],
            ["myAnalyses", "Query", "Autenticado", "Lista el historial del usuario actual."],
            ["listAnalyses", "Query", "Administrador", "Lista global de análisis."],
            ["getSystemStats", "Query", "Administrador", "Devuelve métricas globales del sistema."],
            ["me", "Query", "Autenticado", "Devuelve el usuario autenticado."],
        ],
    )

    heading(doc, "5.3. Modelos de datos persistidos", 2)
    table(
        doc,
        ["Entidad", "Campos principales", "Persistencia"],
        [
            ["UserRecord", "user_id, email, name, password_hash, role, is_active, created_at, updated_at", "Colección `users`."],
            ["AnalysisRecord", "analysis_id, user_id, file_name, file_path, mime_type, detections, status, inference_time_ms, model_version, error_message", "Colección `analyses`."],
            ["DetectionRecord", "class_id, class_name, confidence, bbox_xyxy, label", "Anidado dentro de `AnalysisRecord`."],
        ],
    )

    heading(doc, "5.4. Colecciones MongoDB", 2)
    table(
        doc,
        ["Colección", "Uso"],
        [
            ["users", "Usuarios autenticados del sistema."],
            ["analyses", "Resultados de los análisis radiográficos."],
            ["dead_letter_events", "Eventos que no se pudieron publicar correctamente."],
        ],
    )

    heading(doc, "5.5. Configuración principal del backend", 2)
    table(
        doc,
        ["Variable / clave", "Valor o uso por defecto"],
        [
            ["DENTAL_AI_MONGO_URI", "MongoDB local por defecto: `mongodb://localhost:27017`."],
            ["DENTAL_AI_MONGO_DB_NAME", "Base de datos: `dental_ai`."],
            ["DENTAL_AI_MODEL_PATH", "Ruta al archivo `best.pt` del modelo entrenado."],
            ["DENTAL_AI_MODEL_CONFIDENCE", "Umbral de confianza de inferencia; por defecto 0.25."],
            ["DENTAL_AI_MODEL_IOU", "Umbral IoU; por defecto 0.45."],
            ["DENTAL_AI_UPLOADS_DIR", "Directorio de archivos subidos; por defecto `./storage/uploads`."],
            ["DENTAL_AI_EVENTS_ENABLED", "Habilita o deshabilita el sistema de eventos."],
            ["DENTAL_AI_EVENTS_TRANSPORT", "Transporte de eventos: `log` o `kafka`."],
            ["DENTAL_AI_AUTH_JWT_SECRET", "Clave de firma JWT."],
        ],
    )

    heading(doc, "5.6. Eventos y topics", 2)
    table(
        doc,
        ["Evento", "Topic Kafka", "Propósito"],
        [
            ["image.uploaded", "dental.image.uploaded", "Se publica al guardar una imagen."],
            ["analysis.requested", "dental.analysis.requested", "Solicitud de análisis encolada."],
            ["analysis.started", "dental.analysis.started", "Marca el inicio de la inferencia."],
            ["analysis.completed", "dental.analysis.completed", "Marca análisis exitoso."],
            ["analysis.failed", "dental.analysis.failed", "Marca análisis fallido."],
            ["result.saved", "dental.result.saved", "Confirma persistencia del resultado."],
            ["metrics.inference", "dental.metrics.inference", "Publica métricas de tiempo y detecciones."],
            ["auth.login_success / auth.login_failed / auth.register / auth.logout / auth.token_refresh", "dental.auth.events", "Eventos de autenticación."],
            ["system.log", "dental.system.logs", "Logs estructurados."],
            ["system.error", "dental.system.errors", "Errores estructurados."],
            ["cualquier evento sin mapeo", "dental.dead-letter", "Ruta de contingencia."],
        ],
    )

    heading(doc, "5.7. Flujo de análisis", 2)
    paragraph(doc, "El flujo de análisis es el siguiente:")
    numbered(doc, "El frontend envía la radiografía en base64 a la mutación `uploadRadiography`.")
    numbered(doc, "El backend valida tipo MIME y tamaño, convierte la imagen y la guarda en `storage/uploads`.")
    numbered(doc, "Se publica el evento `image.uploaded` y luego `analysis.requested`.")
    numbered(doc, "El servicio de inferencia carga el modelo YOLO, ejecuta `predict` y normaliza detecciones.")
    numbered(doc, "Se guarda el `AnalysisRecord` en MongoDB con detecciones, tiempos y versión del modelo.")
    numbered(doc, "Se publica el resultado y el frontend redirige al diagnóstico.")

    heading(doc, "5.8. Dependencias y pruebas del backend", 2)
    table(
        doc,
        ["Archivo", "Contenido relevante"],
        [
            ["backend/pyproject.toml", "Define FastAPI, Strawberry GraphQL, Motor, Ultralytics, Pillow, JWT y tests."],
            ["backend/requirements.txt", "Lista de dependencias instalables con pip."],
            ["backend/tests/", "Pruebas smoke de auth, eventos, GraphQL, análisis y salud del sistema."],
        ],
    )
    paragraph(doc, "Comandos habituales de backend:")
    code_block(
        doc,
        "cd backend\npython -m venv .venv\n.venv\\Scripts\\activate\npip install -r requirements.txt\nuvicorn main:app --reload --host 0.0.0.0 --port 8000",
    )

    heading(doc, "6. Frontend", 1)
    paragraph(doc, "El frontend está construido con Vue 3, Vite y Tailwind CSS v4 con daisyUI. Se organiza como una SPA con rutas públicas y protegidas por sesión en `sessionStorage`.")

    heading(doc, "6.1. Rutas principales", 2)
    table(
        doc,
        ["Ruta", "Acceso", "Pantalla"],
        [
            ["/", "Pública", "LandingView"],
            ["/login", "Pública", "LoginView"],
            ["/register", "Pública", "RegisterView"],
            ["/dashboard", "Protegida", "DashboardView"],
            ["/analyze", "Protegida", "AnalyzeView"],
            ["/diagnostic", "Protegida", "DiagnosticView"],
        ],
    )

    heading(doc, "6.2. Componentes y módulos principales", 2)
    table(
        doc,
        ["Módulo", "Responsabilidad"],
        [
            ["DashboardView.vue", "Compone el resumen clínico, el hero de salud, el último diagnóstico y la gráfica de problemas."],
            ["AnalyzeView.vue", "Pantalla de carga y validación de radiografías."],
            ["DiagnosticView.vue", "Pantalla de resultados con imagen anotada y panel de hallazgos."],
            ["HealthScoreHero.vue + useHealthScoreHero.js", "Calculan el promedio de aciertos y la severidad agregada de los hallazgos."],
            ["LatestDiagnosisSummary.vue + useLatestDiagnosisSummary.js", "Resumen del último análisis completado, incluyendo imagen y métricas."],
            ["ProblemTypesChart.vue + useProblemTypesChart.js", "Calcula conteos y porcentajes de caries, empastes, implantes e impactados."],
            ["useUploadImagesQueue.js", "Gestiona la cola de carga, validación, base64 y envío a GraphQL."],
            ["useDiagnosticAnalysis.js", "Persiste el diagnóstico actual en `sessionStorage` y mantiene la imagen seleccionada."],
            ["authService.js", "Gestión de sesión, login, registro y logout."],
            ["graphqlClient.js", "Cliente HTTP GraphQL con soporte de token Bearer."],
        ],
    )

    heading(doc, "6.3. Estado, sesión y persistencia", 2)
    table(
        doc,
        ["Elemento", "Uso"],
        [
            ["sessionStorage.accessToken", "Autoriza peticiones GraphQL."],
            ["sessionStorage.userId", "Identificador del usuario."],
            ["sessionStorage.email", "Correo mostrado en el informe diagnóstico."],
            ["sessionStorage.role", "Permite distinguir usuario y administrador."],
            ["diagnostic.analysis.v1", "Persistencia del diagnóstico actual entre recargas."],
        ],
    )

    heading(doc, "6.4. Responsividad", 2)
    paragraph(doc, "El dashboard se construye con Flexbox y bloques apilables para adaptarse a móvil, tablet y escritorio. No se usa CSS Grid en el layout principal del dashboard.")
    bullet(doc, "En móvil, los módulos se apilan verticalmente.")
    bullet(doc, "En escritorio, los paneles principales se distribuyen en columnas proporcionales.")
    bullet(doc, "El objetivo es evitar dependencias rígidas de grillas y simplificar el reordenamiento de componentes.")

    heading(doc, "6.5. Configuración del frontend", 2)
    table(
        doc,
        ["Archivo", "Detalle"],
        [
            ["frontend/package.json", "Scripts: `dev`, `build`, `preview` y `format`."],
            ["frontend/src/config/graphql.js", "Lee `VITE_GRAPHQL_ENDPOINT` y usa fallback local."],
            ["frontend/src/router/index.js", "Guarda rutas públicas/protegidas y guards de autenticación."],
            ["frontend/.env", "Define la URL del backend GraphQL para Vite."],
        ],
    )
    paragraph(doc, "Comandos habituales de frontend:")
    code_block(doc, "cd frontend\npnpm install\npnpm dev\npnpm build")

    heading(doc, "7. Entrenamiento e inferencia del modelo", 1)
    paragraph(doc, "La carpeta `entrenamiento ia/` contiene el flujo de preparación del dataset, entrenamiento, evaluación y predicción con YOLO. También incluye scripts de validación y pruebas automatizadas.")

    heading(doc, "7.1. Estructura y scripts", 2)
    table(
        doc,
        ["Elemento", "Función"],
        [
            ["tools/csv_to_yolo.py", "Convierte anotaciones CSV al formato YOLO."],
            ["tools/train_yolov8.py", "Entrena un modelo YOLO con parámetros configurables."],
            ["tools/eval_predict_yolov8.py", "Evalúa y/o predice usando un modelo entrenado."],
            ["tools/run_train_eval_predict_yolov8.py", "Pipeline completo de entrenamiento, validación y predicción."],
            ["dataset/", "Imágenes, etiquetas y configuración del dataset."],
            ["runs/", "Resultados de entrenamiento, evaluación y predicción."],
            ["tests/", "Pruebas smoke y de extracción de métricas."],
            ["yolo26*.pt", "Pesos base locales para entrenamiento o fine tuning."],
        ],
    )

    heading(doc, "7.2. Evidencias visuales", 2)
    add_image(doc, IMG_DIR / "captura entrenamiento.png", "Figura 1. Entrenamiento del modelo YOLO en el entorno de IA dental.")
    add_image(doc, IMG_DIR / "validacion.png", "Figura 2. Validación del modelo con métricas y resultados intermedios.")
    add_image(doc, IMG_DIR / "prediccion.png", "Figura 3. Predicción sobre radiografía y visualización de resultados.")

    heading(doc, "7.3. Salidas esperadas del entrenamiento", 2)
    table(
        doc,
        ["Salida", "Ubicación"],
        [
            ["Peso mejor del modelo", "runs/train/<run>/weights/best.pt"],
            ["Último peso del modelo", "runs/train/<run>/weights/last.pt"],
            ["Reporte de evaluación", "runs/eval_predict/<run>/val_metrics.json y val_metrics.csv"],
            ["Resultados de predicción", "runs/eval_predict/<run>/"],
            ["Reporte del pipeline", "runs/pipeline/<run>/pipeline_report.json"],
        ],
    )

    heading(doc, "8. Pruebas y verificación", 1)
    table(
        doc,
        ["Área", "Tipo de prueba"],
        [
            ["backend/tests/", "Smoke tests de auth, eventos, GraphQL y análisis."],
            ["entrenamiento ia/tests/", "Smoke tests del conversor, entrenamiento y pipeline."],
            ["frontend build", "Verificación de compilación con Vite y revisión de errores de plantilla."],
        ],
    )
    paragraph(doc, "Sugerencia de verificación rápida:")
    code_block(doc, "cd backend\npython -m pytest\ncd ..\\frontend\npnpm build\ncd ..\\entrenamiento ia\npython -m pytest")

    heading(doc, "9. Problemas conocidos y recomendaciones", 1)
    table(
        doc,
        ["Problema", "Causa probable", "Solución"],
        [
            ["El backend no arranca", "La ruta de `MODEL_PATH` apunta a un `best.pt` inexistente.", "Corregir la ruta o restaurar el peso entrenado."],
            ["El frontend no ve GraphQL", "URL incorrecta en `VITE_GRAPHQL_ENDPOINT`.", "Ajustar `.env` y reiniciar Vite."],
            ["MongoDB no conecta", "El servicio no está levantado o la URI es incorrecta.", "Iniciar MongoDB y revisar la variable `DENTAL_AI_MONGO_URI`."],
            ["No aparecen resultados en dashboard", "Aún no hay análisis completados en la sesión o en la base de datos.", "Subir una radiografía o revisar el historial del usuario."],
            ["Kafka falla al iniciar", "No está disponible o no hay conexión.", "Usar transporte `log` como fallback o deshabilitar eventos."],
        ],
    )
    paragraph(doc, "Nota importante: en la sesión de trabajo se detectó un fallo de arranque del backend por falta del archivo `best.pt` en la ruta configurada. Ese punto debe resolverse para ejecutar inferencia real.")

    heading(doc, "10. Conclusión", 1)
    paragraph(doc, "El proyecto combina una interfaz web moderna, una API GraphQL, persistencia en MongoDB, un motor de inferencia YOLO y una capa de eventos preparada para crecimiento. Esta documentación resume la arquitectura real del repositorio y sirve como base para despliegue, mantenimiento y futura evolución.")

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()

