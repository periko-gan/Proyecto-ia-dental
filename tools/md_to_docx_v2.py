from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

MD_PATH = Path(__file__).resolve().parents[1] / 'DOCUMENTACION_PROYECTO.md'
OUT_DOCX = Path(__file__).resolve().parents[1] / 'DOCUMENTACION_PROYECTO.docx'


def insert_toc(document: Document):
    p = document.add_paragraph()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(fld)


def add_code_paragraph(document: Document, code_lines: list[str]):
    p = document.add_paragraph()
    run = p.add_run('\n'.join(code_lines))
    run.font.name = 'Courier New'
    run.font.size = Pt(9)


def add_table_from_md(document: Document, table_lines: list[str]):
    def split_row(line: str) -> list[str]:
        parts = [c.strip() for c in re.split(r"\s*\|\s*", line.strip())]
        if parts and parts[0] == '':
            parts = parts[1:]
        if parts and parts[-1] == '':
            parts = parts[:-1]
        return parts

    if not table_lines:
        return
    # detect header and separator
    if len(table_lines) >= 2 and re.match(r"\s*\|?\s*-{3,}", table_lines[1]):
        header = split_row(table_lines[0])
        body = table_lines[2:]
    else:
        # fallback: first line header, rest body
        header = split_row(table_lines[0])
        body = table_lines[1:]

    rows = [split_row(ln) for ln in body if ln.strip() != '']
    cols = max(len(header), max((len(r) for r in rows), default=0))
    table = document.add_table(rows=1 + len(rows), cols=cols)
    table.style = 'Table Grid'
    # header
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
    # body
    for r_idx, row in enumerate(rows, start=1):
        for c_idx in range(cols):
            txt = row[c_idx] if c_idx < len(row) else ''
            table.rows[r_idx].cells[c_idx].text = txt
    document.add_paragraph('')


def md_to_docx(md_path: Path, out_path: Path) -> None:
    with md_path.open(encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    first_heading_seen = False
    in_code = False
    code_buffer: list[str] = []

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip('\n')

        if line.strip().startswith('```'):
            in_code = not in_code
            if not in_code and code_buffer:
                add_code_paragraph(doc, code_buffer)
                code_buffer = []
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # Heading
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not first_heading_seen:
                doc.add_heading(text, level=1)
                insert_toc(doc)
                first_heading_seen = True
            else:
                doc.add_heading(text, level=level)
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^---+$', line.strip()):
            doc.add_page_break()
            i += 1
            continue

        # Unordered list
        mul = re.match(r'^\s*[-\*+]\s+(.*)', line)
        if mul:
            doc.add_paragraph(mul.group(1), style='List Bullet')
            i += 1
            continue

        # Ordered list
        mold = re.match(r'^\s*\d+[\.)]\s+(.*)', line)
        if mold:
            doc.add_paragraph(mold.group(1), style='List Number')
            i += 1
            continue

        # Table detection: a block of lines containing '|' and at least one separator line
        if '|' in line:
            # try to collect a table block
            tbl_lines = [line]
            j = i + 1
            has_sep = False
            while j < len(lines):
                nxt = lines[j].rstrip('\n')
                if nxt.strip() == '':
                    break
                if re.match(r"\s*\|?\s*-{3,}", nxt):
                    has_sep = True
                if '|' not in nxt and not re.match(r"\s*\|?\s*-{3,}", nxt):
                    break
                tbl_lines.append(nxt)
                j += 1
            if len(tbl_lines) > 1 and has_sep:
                add_table_from_md(doc, tbl_lines)
                i = j
                continue
            # else fallthrough to normal paragraph

        # Empty line
        if line.strip() == '':
            doc.add_paragraph('')
            i += 1
            continue

        # Normal paragraph
        doc.add_paragraph(line)
        i += 1

    if code_buffer:
        add_code_paragraph(doc, code_buffer)

    doc.save(out_path)
    print(f'Wrote: {out_path}')


if __name__ == '__main__':
    if not MD_PATH.exists():
        print('Markdown source not found:', MD_PATH)
    else:
        md_to_docx(MD_PATH, OUT_DOCX)

